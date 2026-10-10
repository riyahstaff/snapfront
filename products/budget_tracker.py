"""Budget & Savings Tracker — Excel / Google Sheets template."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from common import OUT

F = "Arial"
INK, SOFT = "2B2A28", "6E6A63"
TC, SAGE, SAND, PAPER = "C8664B", "8BA68A", "E7C9A0", "F7F3EC"
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")
HEAD_FILL = PatternFill("solid", fgColor=INK)
SUB_FILL = PatternFill("solid", fgColor="EFE8DC")
thin = Side(style="thin", color="DDD5C7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '$#,##0.00;($#,##0.00);"-"'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

wb = Workbook()

def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(name=F, bold=True, size=18, color=INK)
    ws["A2"] = sub; ws["A2"].font = Font(name=F, size=10, color=SOFT)
    ws.sheet_view.showGridLines = False

# ---------------- Start Here ----------------
ws = wb.active; ws.title = "Start Here"
title(ws, "Budget & Savings Tracker", "Snapfront Studio · works in Excel, Google Sheets and Numbers")
steps = [
    ("1", "Setup tab", "Enter the year and starting balance, then edit the category list and monthly budget for each category. Yellow cells are yours to change."),
    ("2", "Transactions tab", "Log every income or expense: date, description, pick a category from the dropdown, and the amount (always a positive number)."),
    ("3", "Monthly Summary tab", "Fills in automatically: actual spending by category and month, budget vs. actual, net savings and savings rate."),
    ("4", "Savings Goals tab", "Add up to 10 goals. The sheet tells you how much to save each month to hit the target date."),
    ("5", "Dashboard tab", "Chart of income vs. expenses by month, updated as you log transactions."),
]
ws["A4"] = "How to use it"; ws["A4"].font = Font(name=F, bold=True, size=12, color=TC)
for i, (n, t, d) in enumerate(steps):
    r = 5 + i
    ws.cell(r, 1, n).font = Font(name=F, bold=True, size=11, color=TC)
    ws.cell(r, 2, t).font = Font(name=F, bold=True, size=10)
    ws.cell(r, 3, d).font = Font(name=F, size=10)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32
ws["A11"] = "Legend"; ws["A11"].font = Font(name=F, bold=True, size=12, color=TC)
ws["B12"] = "Input"; ws["B12"].fill = INPUT_FILL; ws["B12"].font = Font(name=F, color="0000FF", size=10)
ws["C12"] = "Yellow cells with blue text are where you type. Everything else is calculated — no need to touch it."
ws["B13"] = "Formula"; ws["B13"].font = Font(name=F, size=10)
ws["C13"] = "Black text cells are formulas. Overwriting them will break the summary."
ws["B14"] = "Example"; ws["B14"].font = Font(name=F, size=10, italic=True, color=SOFT)
ws["C14"] = "Row 5 on the Transactions tab is an example entry — overwrite or delete it when you start."
for r in (12, 13, 14): ws.cell(r, 3).font = Font(name=F, size=10)
ws["A16"] = "Personal use licence. Please do not resell or share this file. © Snapfront Studio."
ws["A16"].font = Font(name=F, size=9, color=SOFT, italic=True)
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 20; ws.column_dimensions["C"].width = 95

# ---------------- Setup ----------------
st = wb.create_sheet("Setup")
title(st, "Setup", "Edit the yellow cells. Categories feed the dropdowns and the monthly summary.")
st["A4"] = "Tracking year"; st["B4"] = 2027
st["A5"] = "Starting balance"; st["B5"] = 0
for c in ("B4", "B5"):
    st[c].fill = INPUT_FILL; st[c].font = Font(name=F, color="0000FF", bold=True); st[c].border = BOX
st["B5"].number_format = MONEY
st["A4"].font = st["A5"].font = Font(name=F, bold=True, size=10)
st["C4"] = "Transactions outside this year are ignored by the summary."; st["C4"].font = Font(name=F, size=9, color=SOFT)

st["A7"], st["B7"], st["C7"] = "Category", "Type", "Monthly budget"
style_header(st, 7, 1, 3)
CATS = [("Salary / wages", "Income", 4200), ("Side hustle", "Income", 600), ("Other income", "Income", 0),
        ("Rent / mortgage", "Expense", 1500), ("Utilities", "Expense", 180), ("Groceries", "Expense", 500),
        ("Transport", "Expense", 250), ("Insurance", "Expense", 150), ("Phone & internet", "Expense", 110),
        ("Subscriptions", "Expense", 45), ("Eating out", "Expense", 200), ("Health", "Expense", 80),
        ("Personal care", "Expense", 60), ("Clothing", "Expense", 80), ("Entertainment", "Expense", 100),
        ("Gifts & giving", "Expense", 75), ("Debt payments", "Expense", 300), ("Savings transfer", "Expense", 500),
        ("Miscellaneous", "Expense", 100)]
CAT_FIRST, CAT_LAST = 8, 8 + 29   # 30 category slots
for i in range(30):
    r = CAT_FIRST + i
    if i < len(CATS):
        st.cell(r, 1, CATS[i][0]); st.cell(r, 2, CATS[i][1]); st.cell(r, 3, CATS[i][2])
    for c in (1, 2, 3):
        cell = st.cell(r, c); cell.fill = INPUT_FILL; cell.font = Font(name=F, color="0000FF", size=10); cell.border = BOX
    st.cell(r, 3).number_format = MONEY
type_dv = DataValidation(type="list", formula1='"Income,Expense"', allow_blank=True)
st.add_data_validation(type_dv); type_dv.add(f"B{CAT_FIRST}:B{CAT_LAST}")
st["C7"].comment = Comment("Example amounts — replace with your own budget.", "Snapfront")
st["E7"] = "Budget totals"; st["E7"].font = Font(name=F, bold=True, size=11, color=TC)
st["E8"], st["F8"] = "Monthly income budget", f'=SUMIFS(C{CAT_FIRST}:C{CAT_LAST},B{CAT_FIRST}:B{CAT_LAST},"Income")'
st["E9"], st["F9"] = "Monthly expense budget", f'=SUMIFS(C{CAT_FIRST}:C{CAT_LAST},B{CAT_FIRST}:B{CAT_LAST},"Expense")'
st["E10"], st["F10"] = "Planned left over", "=F8-F9"
for r in (8, 9, 10):
    st[f"E{r}"].font = Font(name=F, size=10); st[f"F{r}"].font = Font(name=F, size=10, bold=(r == 10)); st[f"F{r}"].number_format = MONEY
st.column_dimensions["A"].width = 24; st.column_dimensions["B"].width = 12; st.column_dimensions["C"].width = 16
st.column_dimensions["D"].width = 4; st.column_dimensions["E"].width = 24; st.column_dimensions["F"].width = 16
CAT_RANGE = f"Setup!$A${CAT_FIRST}:$A${CAT_LAST}"
TYPE_RANGE = f"Setup!$B${CAT_FIRST}:$B${CAT_LAST}"

# ---------------- Transactions ----------------
tx = wb.create_sheet("Transactions")
title(tx, "Transactions", "One row per transaction. Amounts are always positive — the category decides income vs. expense.")
heads = ["Date", "Description", "Category", "Amount", "Type", "Month", "Year", "Notes"]
for i, h in enumerate(heads, 1): tx.cell(4, i, h)
style_header(tx, 4, 1, 8)
TX_FIRST, TX_LAST = 5, 1004
for r in range(TX_FIRST, TX_LAST + 1):
    for c in (1, 2, 3, 4, 8):
        cell = tx.cell(r, c); cell.fill = INPUT_FILL; cell.font = Font(name=F, color="0000FF", size=10); cell.border = BOX
    tx.cell(r, 1).number_format = "mm/dd/yyyy"
    tx.cell(r, 4).number_format = MONEY
    tx.cell(r, 5, f'=IF(C{r}="","",IFERROR(INDEX({TYPE_RANGE},MATCH(C{r},{CAT_RANGE},0)),"Check category"))')
    tx.cell(r, 6, f'=IF(A{r}="","",MONTH(A{r}))')
    tx.cell(r, 7, f'=IF(A{r}="","",YEAR(A{r}))')
    for c in (5, 6, 7):
        tx.cell(r, c).font = Font(name=F, size=10); tx.cell(r, c).border = BOX
        tx.cell(r, c).alignment = Alignment(horizontal="center")
import datetime as dt
tx["A5"] = dt.date(2027, 1, 3); tx["B5"] = "Example: weekly grocery shop"; tx["C5"] = "Groceries"; tx["D5"] = 86.40
tx["H5"] = "Example row — overwrite me"
for c in ("A5", "B5", "C5", "D5", "H5"): tx[c].font = Font(name=F, color="0000FF", size=10, italic=True)
cat_dv = DataValidation(type="list", formula1=CAT_RANGE, allow_blank=True, showErrorMessage=True,
                        errorTitle="Unknown category", error="Pick a category from the list, or add it on the Setup tab.")
tx.add_data_validation(cat_dv); cat_dv.add(f"C{TX_FIRST}:C{TX_LAST}")
amt_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, error="Enter a positive amount.")
tx.add_data_validation(amt_dv); amt_dv.add(f"D{TX_FIRST}:D{TX_LAST}")
tx.freeze_panes = "A5"
for col, w in zip("ABCDEFGH", [13, 34, 22, 14, 11, 8, 8, 30]): tx.column_dimensions[col].width = w
tx.auto_filter.ref = f"A4:H{TX_LAST}"

# ---------------- Monthly Summary ----------------
sm = wb.create_sheet("Monthly Summary")
title(sm, "Monthly Summary", "Calculated automatically from Transactions for the year on the Setup tab.")
sm["A4"] = "Category"; sm["B4"] = "Type"
for i, m in enumerate(MONTHS): sm.cell(4, 3 + i, m)
sm.cell(4, 15, "Year total"); sm.cell(4, 16, "Year budget"); sm.cell(4, 17, "Difference")
style_header(sm, 4, 1, 17)
A, C, D, E, F_, G = (f"Transactions!$A${TX_FIRST}:$A${TX_LAST}", f"Transactions!$C${TX_FIRST}:$C${TX_LAST}",
                     f"Transactions!$D${TX_FIRST}:$D${TX_LAST}", f"Transactions!$E${TX_FIRST}:$E${TX_LAST}",
                     f"Transactions!$F${TX_FIRST}:$F${TX_LAST}", f"Transactions!$G${TX_FIRST}:$G${TX_LAST}")
SUM_FIRST = 5
for i in range(30):
    r = SUM_FIRST + i; sr = CAT_FIRST + i
    sm.cell(r, 1, f'=IF(Setup!A{sr}="","",Setup!A{sr})')
    sm.cell(r, 2, f'=IF(Setup!B{sr}="","",Setup!B{sr})')
    for m in range(12):
        col = 3 + m
        sm.cell(r, col, f'=IF($A{r}="",0,SUMIFS({D},{C},$A{r},{F_},{m+1},{G},Setup!$B$4))')
        sm.cell(r, col).number_format = MONEY
    sm.cell(r, 15, f"=SUM(C{r}:N{r})")
    sm.cell(r, 16, f"=IF($A{r}=\"\",0,Setup!C{sr}*12)")
    sm.cell(r, 17, f'=IF($A{r}="",0,IF(B{r}="Income",O{r}-P{r},P{r}-O{r}))')
    for c in range(1, 18):
        cell = sm.cell(r, c); cell.font = Font(name=F, size=9.5); cell.border = BOX
        if c >= 15: cell.number_format = MONEY
    sm.cell(r, 15).font = Font(name=F, size=9.5, bold=True)
SUM_LAST = SUM_FIRST + 29
tr = SUM_LAST + 2
labels = [("Total income", "Income"), ("Total expenses", "Expense")]
for k, (lab, typ) in enumerate(labels):
    r = tr + k
    sm.cell(r, 1, lab)
    for c in range(3, 18):
        L = sm.cell(4, c).column_letter
        sm.cell(r, c, f'=SUMIFS({L}${SUM_FIRST}:{L}${SUM_LAST},$B${SUM_FIRST}:$B${SUM_LAST},"{typ}")')
r_net = tr + 2; r_rate = tr + 3; r_bal = tr + 4
sm.cell(r_net, 1, "Net (income − expenses)")
sm.cell(r_rate, 1, "Savings rate")
sm.cell(r_bal, 1, "Running balance")
for c in range(3, 16):
    L = sm.cell(4, c).column_letter
    sm.cell(r_net, c, f"={L}{tr}-{L}{tr+1}")
    sm.cell(r_rate, c, f"=IF({L}{tr}=0,0,{L}{r_net}/{L}{tr})")
for c in range(3, 15):
    L = sm.cell(4, c).column_letter
    prev = "Setup!$B$5" if c == 3 else f"{sm.cell(4, c-1).column_letter}{r_bal}"
    sm.cell(r_bal, c, f"={prev}+{L}{r_net}")
for r in range(tr, r_bal + 1):
    sm.cell(r, 1).font = Font(name=F, bold=True, size=10)
    for c in range(1, 18):
        cell = sm.cell(r, c); cell.fill = SUB_FILL; cell.border = BOX
        if c >= 3:
            cell.number_format = "0.0%" if r == r_rate else MONEY
            cell.font = Font(name=F, size=9.5, bold=True)
sm.conditional_formatting.add(f"Q{SUM_FIRST}:Q{SUM_LAST}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C0392B")))
sm.cell(r_bal + 2, 1, "Difference: positive = better than budget (earned more or spent less). Red = over budget.").font = Font(name=F, size=9, italic=True, color=SOFT)
sm.column_dimensions["A"].width = 24; sm.column_dimensions["B"].width = 9
for c in range(3, 18): sm.column_dimensions[sm.cell(4, c).column_letter].width = 11.5
sm.freeze_panes = "C5"

# ---------------- Savings Goals ----------------
sg = wb.create_sheet("Savings Goals")
title(sg, "Savings Goals", "Add a goal, a target and a date. The sheet works out what to save each month.")
for i, h in enumerate(["Goal", "Target amount", "Saved so far", "Remaining", "Progress", "Target date", "Months left", "Save per month"], 1):
    sg.cell(4, i, h)
style_header(sg, 4, 1, 8)
for i in range(10):
    r = 5 + i
    for c in (1, 2, 3, 6):
        cell = sg.cell(r, c); cell.fill = INPUT_FILL; cell.font = Font(name=F, color="0000FF", size=10)
    sg.cell(r, 4, f'=IF(B{r}="","",MAX(B{r}-C{r},0))')
    sg.cell(r, 5, f'=IF(OR(B{r}="",B{r}=0),"",MIN(C{r}/B{r},1))')
    sg.cell(r, 7, f'=IF(F{r}="","",MAX(0,(YEAR(F{r})-YEAR(TODAY()))*12+MONTH(F{r})-MONTH(TODAY())))')
    sg.cell(r, 8, f'=IF(OR(D{r}="",G{r}=""),"",IF(G{r}=0,D{r},D{r}/G{r}))')
    for c in range(1, 9):
        sg.cell(r, c).border = BOX
        if c in (4, 5, 7, 8): sg.cell(r, c).font = Font(name=F, size=10)
    for c in (2, 3, 4, 8): sg.cell(r, c).number_format = MONEY
    sg.cell(r, 5).number_format = "0%"; sg.cell(r, 6).number_format = "mm/dd/yyyy"
    sg.cell(r, 7).alignment = Alignment(horizontal="center")
sg["A5"], sg["B5"], sg["C5"], sg["F5"] = "Example: emergency fund", 3000, 450, dt.date(2027, 12, 31)
sg["A5"].font = Font(name=F, color="0000FF", size=10, italic=True)
sg.conditional_formatting.add("E5:E14", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=SAGE))
for col, w in zip("ABCDEFGH", [28, 15, 15, 14, 12, 14, 12, 15]): sg.column_dimensions[col].width = w

# ---------------- Dashboard ----------------
db = wb.create_sheet("Dashboard")
title(db, "Dashboard", "Income vs. expenses by month — updates as you log transactions.")
db["A4"] = "Month"; db["B4"] = "Income"; db["C4"] = "Expenses"; db["D4"] = "Net"
style_header(db, 4, 1, 4)
for m in range(12):
    r = 5 + m; L = sm.cell(4, 3 + m).column_letter
    db.cell(r, 1, MONTHS[m])
    db.cell(r, 2, f"='Monthly Summary'!{L}{tr}")
    db.cell(r, 3, f"='Monthly Summary'!{L}{tr+1}")
    db.cell(r, 4, f"=B{r}-C{r}")
    for c in range(1, 5):
        db.cell(r, c).border = BOX; db.cell(r, c).font = Font(name=F, size=10, color="008000" if c in (2, 3) else "000000")
        if c > 1: db.cell(r, c).number_format = MONEY
db["A18"] = "Year"; db["B18"] = "=SUM(B5:B16)"; db["C18"] = "=SUM(C5:C16)"; db["D18"] = "=B18-C18"
for c in "ABCD":
    db[f"{c}18"].font = Font(name=F, bold=True, size=10); db[f"{c}18"].fill = SUB_FILL; db[f"{c}18"].border = BOX
    if c != "A": db[f"{c}18"].number_format = MONEY
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.title = "Income vs. expenses"
ch.y_axis.title = "Amount ($)"; ch.y_axis.numFmt = "$#,##0"; ch.height = 9; ch.width = 20
ch.add_data(Reference(db, min_col=2, min_row=4, max_col=3, max_row=16), titles_from_data=True)
ch.set_categories(Reference(db, min_col=1, min_row=5, max_row=16))
ch.series[0].graphicalProperties.solidFill = SAGE; ch.series[1].graphicalProperties.solidFill = TC
ch.y_axis.delete = False; ch.x_axis.delete = False
db.add_chart(ch, "F4")
for col, w in zip("ABCD", [10, 14, 14, 14]): db.column_dimensions[col].width = w

for sh in wb.worksheets:
    sh.sheet_properties.tabColor = {"Start Here": TC, "Setup": SAND, "Transactions": SAGE}.get(sh.title, "7D8FB3")
path = os.path.join(OUT, "Snapfront-Budget-Savings-Tracker.xlsx")
wb.save(path); print(path)
